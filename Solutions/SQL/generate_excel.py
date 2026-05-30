"""
Generates sample_data.xlsx for the ANSI SQL Module 1 exercise database.
Requires: pip install openpyxl
Run:      python generate_excel.py
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FONT  = Font(bold=True, color="FFFFFF")
HEADER_FILL  = PatternFill(fill_type="solid", fgColor="1F3864")
HEADER_ALIGN = Alignment(horizontal="center")

SHEETS = {
    "Users": {
        "headers": ["user_id", "full_name", "email", "city", "registration_date"],
        "rows": [
            (1, "Alice Johnson", "alice@example.com",   "New York",    "2024-12-01"),
            (2, "Bob Smith",     "bob@example.com",     "Los Angeles", "2024-12-05"),
            (3, "Charlie Lee",   "charlie@example.com", "Chicago",     "2024-12-10"),
            (4, "Diana King",    "diana@example.com",   "New York",    "2025-01-15"),
            (5, "Ethan Hunt",    "ethan@example.com",   "Los Angeles", "2025-02-01"),
        ],
    },
    "Events": {
        "headers": ["event_id", "title", "description", "city", "start_date", "end_date", "status", "organizer_id"],
        "rows": [
            (1, "Tech Innovators Meetup",        "A meetup for tech enthusiasts.",        "New York",    "2025-06-10 10:00:00", "2025-06-10 16:00:00", "upcoming",  1),
            (2, "AI & ML Conference",            "Conference on AI and ML advancements.", "Chicago",     "2025-05-15 09:00:00", "2025-05-15 17:00:00", "completed", 3),
            (3, "Frontend Development Bootcamp", "Hands-on training on frontend tech.",   "Los Angeles", "2025-07-01 10:00:00", "2025-07-03 16:00:00", "upcoming",  2),
        ],
    },
    "Sessions": {
        "headers": ["session_id", "event_id", "title", "speaker_name", "start_time", "end_time"],
        "rows": [
            (1, 1, "Opening Keynote",   "Dr. Tech",      "2025-06-10 10:00:00", "2025-06-10 11:00:00"),
            (2, 1, "Future of Web Dev", "Alice Johnson",  "2025-06-10 11:15:00", "2025-06-10 12:30:00"),
            (3, 2, "AI in Healthcare",  "Charlie Lee",    "2025-05-15 09:30:00", "2025-05-15 11:00:00"),
            (4, 3, "Intro to HTML5",    "Bob Smith",      "2025-07-01 10:00:00", "2025-07-01 12:00:00"),
        ],
    },
    "Registrations": {
        "headers": ["registration_id", "user_id", "event_id", "registration_date"],
        "rows": [
            (1, 1, 1, "2025-05-01"),
            (2, 2, 1, "2025-05-02"),
            (3, 3, 2, "2025-04-30"),
            (4, 4, 2, "2025-04-28"),
            (5, 5, 3, "2025-06-15"),
        ],
    },
    "Feedback": {
        "headers": ["feedback_id", "user_id", "event_id", "rating", "comments", "feedback_date"],
        "rows": [
            (1, 3, 2, 4, "Great insights!",   "2025-05-16"),
            (2, 4, 2, 5, "Very informative.", "2025-05-16"),
            (3, 2, 1, 3, "Could be better.",  "2025-06-11"),
        ],
    },
    "Resources": {
        "headers": ["resource_id", "event_id", "resource_type", "resource_url", "uploaded_at"],
        "rows": [
            (1, 1, "pdf",   "https://portal.com/resources/tech_meetup_agenda.pdf", "2025-05-01 10:00:00"),
            (2, 2, "image", "https://portal.com/resources/ai_poster.jpg",          "2025-04-20 09:00:00"),
            (3, 3, "link",  "https://portal.com/resources/html5_docs",             "2025-06-25 15:00:00"),
        ],
    },
}


def build_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font  = HEADER_FONT
        cell.fill  = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    for row in rows:
        ws.append(list(row))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4


def main():
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, data in SHEETS.items():
        ws = wb.create_sheet(title=sheet_name)
        build_sheet(ws, data["headers"], data["rows"])

    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_data.xlsx")
    wb.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
